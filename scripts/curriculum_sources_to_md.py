"""Convert the official curriculum PDFs in docs/curriculum_sources to Markdown.

The PDFs carry a text layer, but their Thai decodes badly: InDesign maps the
TH Sarabun glyph for า to ำ and nikhahit to U+FFFD, Word maps nikhahit and the
lowered mai ek to spaces, old files use Thai PUA glyph variants, and the 2551
core curriculum is Thai bytes shown through MacRoman. This script reads each
glyph's id, decodes TH Sarabun glyphs through the installed font's own glyph
names, and repairs the rest by rule. Each page is written under a
``<!-- page N -->`` marker whose N is the PDF page number, and a scored
QUALITY.md is written alongside.

Needs the ``curriculum-sources`` extra and the TH Sarabun PSK font installed
(the glyph repair reads its glyph names; without it the repair is skipped).

    python scripts/curriculum_sources_to_md.py            # every file
    python scripts/curriculum_sources_to_md.py กลุ่มที่1   # names containing this
"""

from __future__ import annotations

import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import pymupdf
from fontTools.ttLib import TTFont

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "docs" / "curriculum_sources"
OUT = SOURCES / "_md"

IN_SCOPE = (
    "_ทั่วไป", "คณิตศาสตร์_ประถมศึกษา", "คณิตศาสตร์_มัธยมศึกษา",
    "วิทยาศาสตร์_ปฐมวัย", "วิทยาศาสตร์_ประถมศึกษา", "วิทยาศาสตร์_มัธยมศึกษา",
    "ฟิสิกส์", "เคมี", "ชีววิทยา", "โลก_ดาราศาสตร์และอวกาศ",
)
# The สพฐ copy of the 2551 curriculum lost every ส to a plain space; the
# สสวท copy of the same book kept it, so only that copy is converted.
SKIP = {"หลักสูตรแกนกลางการศึกษาขั้นพื้นฐาน_2551_ฉบับ_สพฐ.pdf"}

FONT_DIRS = (
    Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "Windows" / "Fonts",
    Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts",
)
SARABUN_FILES = {
    ("psk", False): "THSarabun.ttf", ("psk", True): "THSarabun Bold.ttf",
    ("new", False): "THSarabunNew.ttf", ("new", True): "THSarabunNew Bold.ttf",
}

# Thai presentation variants in the private use area (Microsoft/Adobe layout).
THAI_PUA = {
    0xF700: "\u0e10", 0xF701: "\u0e34", 0xF702: "\u0e35", 0xF703: "\u0e36",
    0xF704: "\u0e37", 0xF705: "\u0e48", 0xF706: "\u0e49", 0xF707: "\u0e4a",
    0xF708: "\u0e4b", 0xF709: "\u0e4c", 0xF70A: "\u0e48", 0xF70B: "\u0e49",
    0xF70C: "\u0e4a", 0xF70D: "\u0e4b", 0xF70E: "\u0e4c", 0xF70F: "\u0e0d",
    0xF710: "\u0e31", 0xF711: "\u0e4d", 0xF712: "\u0e47", 0xF713: "\u0e48",
    0xF714: "\u0e49", 0xF715: "\u0e4a", 0xF716: "\u0e4b", 0xF717: "\u0e4c",
    0xF718: "\u0e38", 0xF719: "\u0e39", 0xF71A: "\u0e3a",
}
# Word writes Symbol-font characters at U+F000 + their Symbol code.
SYMBOL = dict(zip(
    range(0x20, 0x80),
    " !∀#∃%&∋()∗+,−./0123456789:;<=>?≅ΑΒΧΔΕΦΓΗΙϑΚΛΜΝΟΠΘΡΣΤΥςΩΞΨΖ[∴]⊥_"
    "‾αβχδεφγηιϕκλμνοπθρστυϖωξψζ{|}∼ ",
))
SYMBOL.update(zip(
    range(0xA0, 0x100),
    " ϒ′≤⁄∞ƒ♣♦♥♠↔←↑→↓°±″≥×∝∂•÷≠≡≈…⏐⎯↵ℵℑℜ℘⊗⊕∅∩∪⊃⊇⊄⊂⊆∈∉∠∇®©™∏√⋅¬∧∨⇔⇐⇑⇒⇓"
    "◊〈®©™∑⎛⎜⎝⎡⎢⎣⎧⎨⎩⎪ 〉∫⌠⎮⌡⎞⎟⎠⎤⎥⎦⎫⎬⎭ ",
))
BRACKET_PIECES = set("⎛⎜⎝⎡⎢⎣⎧⎨⎩⎪⎞⎟⎠⎤⎥⎦⎫⎬⎭⌠⎮⌡⏐⎯")

TONE = "\u0e48-\u0e4b"
MARKS = "\u0e31\u0e34-\u0e3a\u0e47-\u0e4e"


# --------------------------------------------------------------------------
# Glyph decoding


def _find_font(name: str) -> Path | None:
    for d in FONT_DIRS:
        if (d / name).exists():
            return d / name
    return None


def _glyph_char(glyph_name: str, cmap_rev: dict[str, int]) -> str | None:
    if glyph_name in cmap_rev and not 0xE000 <= cmap_rev[glyph_name] <= 0xF8FF:
        return chr(cmap_rev[glyph_name])
    m = re.match(r"uni([0-9A-F]{4})(?:\..*)?$", glyph_name)
    if m:
        return chr(int(m.group(1), 16))
    if glyph_name in cmap_rev:
        return THAI_PUA.get(cmap_rev[glyph_name])
    return None


def load_sarabun_tables() -> dict[tuple[str, bool], dict[int, str]]:
    tables = {}
    for key, fname in SARABUN_FILES.items():
        path = _find_font(fname)
        if path is None:
            continue
        font = TTFont(path)
        rev: dict[str, int] = {}
        for code, gname in sorted(font.getBestCmap().items()):
            rev.setdefault(gname, code)
        table = {}
        for gid, gname in enumerate(font.getGlyphOrder()):
            ch = _glyph_char(gname, rev)
            if ch is not None:
                table[gid] = ch
        tables[key] = table
    return tables


def sarabun_key(font_name: str) -> tuple[str, bool] | None:
    n = font_name.split("+")[-1].replace(" ", "").lower()
    if not n.startswith("thsarabun"):
        return None
    family = "new" if n.startswith("thsarabunnew") else "psk"
    if n.startswith("thsarabunit"):
        return None
    return family, "bold" in n


class GlyphFixer:
    """Replace a Sarabun character with the one its glyph id really draws.

    A font's table is used on a page only after its glyphs agree with the
    extracted text on most unambiguous consonants, which proves the embedded
    subset kept the original glyph ids.
    """

    def __init__(self) -> None:
        self.tables = load_sarabun_tables()
        self.trusted: dict[tuple[str, str], bool] = {}
        self.stats: Counter = Counter()

    def page_map(self, page: pymupdf.Page, doc_id: str) -> dict:
        by_pos: dict[tuple, list[str]] = defaultdict(list)
        votes: dict[str, list[int]] = defaultdict(lambda: [0, 0])
        spans = []
        for sp in page.get_texttrace():
            key = sarabun_key(sp["font"])
            if key is None or key not in self.tables:
                continue
            spans.append((sp, key))
            table = self.tables[key]
            for c, g, *_ in sp["chars"]:
                ch = chr(c)
                if "\u0e01" <= ch <= "\u0e2e":
                    votes[sp["font"]][0] += table.get(g) == ch
                    votes[sp["font"]][1] += 1
        for font, (ok, n) in votes.items():
            tk = (doc_id, font)
            if n >= 20 or tk not in self.trusted:
                self.trusted[tk] = n == 0 or ok / n >= 0.98
        for sp, key in spans:
            if not self.trusted.get((doc_id, sp["font"]), False):
                continue
            table = self.tables[key]
            for c, g, origin, _bbox in sp["chars"]:
                real = table.get(g)
                # Only Thai glyphs are trusted: some Word subsets renumber
                # the Latin glyphs below them.
                if real is None or not "ก" <= real <= "๛":
                    continue
                pos = (round(origin[0], 1), round(origin[1], 1))
                by_pos[pos].append((chr(c), real))
        return by_pos

    def fix(self, by_pos: dict, char: dict) -> str:
        # Zero-width marks can share an origin with the next letter, so match
        # on the extracted character too. The two extraction paths disagree
        # only on glyphs with no sound mapping (U+FFFD in one, Ğ in the other).
        c = char["c"]
        pos = (round(char["origin"][0], 1), round(char["origin"][1], 1))
        queue = by_pos.get(pos, [])
        pick = next((i for i, (orig, _) in enumerate(queue) if orig == c), None)
        if pick is None and not _reliable(c):
            pick = next((i for i, (orig, _) in enumerate(queue) if not _reliable(orig)), None)
        if pick is None:
            return c
        real = queue.pop(pick)[1]
        if real != c:
            self.stats[(c, real)] += 1
        return real


def _reliable(c: str) -> bool:
    return c.isalnum() and not 0xE000 <= ord(c) <= 0xF8FF and (
        "฀" <= c <= "๿" or c.isascii())


# --------------------------------------------------------------------------
# Text repair


def mac_thai(text: str) -> str:
    """Undo Thai TIS-620 bytes that were decoded as MacRoman."""
    out = []
    for ch in text:
        if 0x80 <= ord(ch) <= 0x9F:
            out.append(SHIFTED_MARKS.get(ord(ch), ch))
            continue
        try:
            b = ch.encode("mac_roman")
            t = b.decode("tis-620") if b[0] >= 0x80 else ch
            out.append(SHIFTED_MARKS.get(ord(t), t))
        except UnicodeError:
            out.append("ต" if ch == "μ" else ch)
    return "".join(out)


# Windows Thai shifted marks at 0x80-0x9F, read from context in the 2551 book.
SHIFTED_MARKS = {
    0x83: "\u0e48", 0x84: "\u0e49", 0x85: "\u0e4a", 0x86: "\u0e4b", 0x87: "\u0e4c",
    0x88: "\u0e48", 0x89: "\u0e49", 0x8A: "\u0e4a", 0x8B: "\u0e4b", 0x8C: "\u0e4c",
    0x92: "\u0e31", 0x93: "\u0e47", 0x94: "\u0e34", 0x95: "\u0e35", 0x96: "\u0e36",
    0x97: "\u0e37", 0x98: "\u0e48", 0x99: "\u0e49", 0x9A: "\u0e4a", 0x9B: "\u0e4b",
    0x9C: "\u0e4c",
}

# Glyphs that heading and IPST fonts map to Latin letters or combining
# accents; each was read from its Thai context. Applied only after Thai text.
THAI_LOOKALIKES = {
    "\u0303": "\u0e49", "\u0306": "\u0e4c", "\u0302": "\u0e48", "\u02e2": "\u0e31",
    "\u02ed": "\u0e47", "\u0114": "\u0e48", "\u0115": "\u0e49", "\u0118": "\u0e4c",
    "\u00ff": "\u0e35", "\u00e9": "\u0e35", "\u00e8": "\u0e34", "\u0119": "\u0e48",
    "\u00ea": "\u0e35", "\u011a": "\u0e49", "\u011d": "\u0e4c", "\u00f5": "\u0e34",
    "\u01e5": "\u0e32",
}
# Leftovers of the nikhahit drawn after ำ.
AM_TAILS = "\ufffd\u01e5"


def looks_mac_thai(text: str) -> bool:
    thai = sum("\u0e00" <= c <= "\u0e7f" for c in text)
    latin1 = sum(c in "“”‘’÷≈—°μ√·π∑»…‡„ß¡¢∫" for c in text)
    return thai < 20 and latin1 > 40


def symbol_chars(text: str) -> str:
    return "".join(SYMBOL.get(ord(c) - 0xF000, c) if 0xF020 <= ord(c) <= 0xF0FF else c
                   for c in text)


def repair(text: str) -> str:
    text ="".join(THAI_PUA.get(ord(c), c) for c in text)
    # ำ is drawn as nikhahit + า, in either order, possibly around a tone mark.
    text = re.sub(f"\u0e4d([{TONE}]?)\u0e32", "\\1\u0e33", text)
    text = re.sub(f"([{TONE}]?)\u0e32([{TONE}]?)\u0e4d", "\\1\\2\u0e33", text)
    text = re.sub(f"\u0e33[{AM_TAILS}]", "\u0e33", text)
    text = re.sub(f"\u0e4d([{TONE}]?)\u0e33", "\1\u0e33", text)
    text = re.sub("(?<=[\u0e01-\u0e4e])[" + "".join(THAI_LOOKALIKES) + "]",
                  lambda m: THAI_LOOKALIKES[m.group()], text)
    text = text.replace("\u2009", " ")
    # Tone mark typed before an upper vowel, and doubled marks.
    text = re.sub(f"([{TONE}\u0e4c])([\u0e31\u0e34-\u0e37\u0e47])", "\\2\\1", text)
    text = re.sub(f"([{MARKS}])\\1+", "\\1", text)
    return unicodedata.normalize("NFC", text)


# --------------------------------------------------------------------------
# Page layout


def line_text(line: dict, fixer: GlyphFixer, by_pos: dict) -> str:
    parts = []
    for span in line["spans"]:
        chars = "".join(fixer.fix(by_pos, ch) for ch in span["chars"])
        if "tflanna" in span["font"].lower():
            # TFLanna labels า as ำ and draws ำ as ė + that glyph.
            chars = re.sub(f"\u0117([{TONE}]?)\u0e33", "\\1\ue000", chars)
            chars = chars.replace("\u0e33", "\u0e32").replace("\ue000", "\u0e33")
        if "symbol" in span["font"].lower() or any(0xF020 <= ord(c) <= 0xF0FF for c in chars):
            chars = symbol_chars(chars)
        parts.append(chars)
    return "".join(parts)


def page_markdown(page: pymupdf.Page, fixer: GlyphFixer, doc_id: str,
                  legacy: bool) -> str:
    by_pos = fixer.page_map(page, doc_id)
    raw = page.get_text("rawdict", flags=pymupdf.TEXTFLAGS_TEXT)
    blocks = []
    sizes = [s["size"] for b in raw["blocks"] for l in b.get("lines", ())
             for s in l["spans"] if s["chars"]]
    body = sorted(sizes)[len(sizes) // 2] if sizes else 0
    for block in raw["blocks"]:
        lines = []
        for line in block.get("lines", ()):
            t = line_text(line, fixer, by_pos)
            t = repair(mac_thai(t) if legacy else t)
            t = "".join(c for c in t if c not in BRACKET_PIECES)
            if t.strip():
                lines.append(t.rstrip())
        if not lines:
            continue
        top = max((s["size"] for l in block["lines"] for s in l["spans"] if s["chars"]),
                  default=0)
        text = "\n".join(lines)
        if body and top >= body * 1.35 and len(text) < 120 and len(lines) <= 2:
            text = "### " + " ".join(l.strip() for l in lines)
        blocks.append(text)
    return "\n\n".join(blocks)


# --------------------------------------------------------------------------
# Drivers


def convert_pdf(path: Path, fixer: GlyphFixer) -> tuple[str, list[str]]:
    doc = pymupdf.open(path)
    doc_id = path.name
    sample = "".join(doc[i].get_text() for i in range(min(8, doc.page_count)))
    legacy = looks_mac_thai(sample)
    pages = []
    for page in doc:
        pages.append(page_markdown(page, fixer, doc_id, legacy))
    return ("legacy-mac-thai" if legacy else "pdf"), pages


def convert_docx(path: Path) -> list[str]:
    import docx

    d = docx.Document(path)
    out = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in block.iter() if n.tag.endswith("}t"))
            if t.strip():
                out.append(repair(t))
        elif tag == "tbl":
            rows = []
            for tr in block.iter():
                if not tr.tag.endswith("}tr"):
                    continue
                cells = []
                for tc in tr:
                    if tc.tag.endswith("}tc"):
                        txt = " ".join(
                            "".join(n.text or "" for n in p.iter() if n.tag.endswith("}t"))
                            for p in tc.iter() if p.tag.endswith("}p"))
                        cells.append(repair(txt).replace("|", "\\|").strip())
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                n = rows[0].count(" | ") + 1
                rows.insert(1, "|" + " --- |" * n)
                out.append("\n".join(rows))
    return ["\n\n".join(out)]


def header(src: Path, kind: str, pages: int) -> str:
    rel = src.relative_to(SOURCES).as_posix()
    return (
        f"---\nsource: {rel}\nkind: {kind}\npages: {pages}\n"
        "generator: scripts/curriculum_sources_to_md.py\n---\n\n"
        f"# {src.stem}\n\n"
        "Machine-extracted from the official PDF. Page markers give the PDF "
        "page number (not the printed folio). Check wording against the PDF "
        "page before quoting it as final.\n"
    )


def main(argv: list[str]) -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    only = argv[1:] or None
    fixer = GlyphFixer()
    if not fixer.tables:
        print("TH Sarabun is not installed; glyph repair disabled", file=sys.stderr)
    for folder in IN_SCOPE:
        for src in sorted((SOURCES / folder).glob("*")):
            if src.name in SKIP or src.suffix.lower() not in (".pdf", ".docx"):
                continue
            if only and not any(o in src.name for o in only):
                continue
            dest = OUT / folder / (src.stem + ".md")
            dest.parent.mkdir(parents=True, exist_ok=True)
            if src.suffix.lower() == ".docx":
                kind, pages = "docx", convert_docx(src)
            else:
                kind, pages = convert_pdf(src, fixer)
            body = [header(src, kind, len(pages))]
            for i, p in enumerate(pages, 1):
                body.append(f"\n<!-- page {i} -->\n\n{p}\n" if kind != "docx" else f"\n{p}\n")
            dest.write_text("".join(body), encoding="utf-8", newline="\n")
            print(f"{kind:16s} {len(pages):4d}p  {dest.relative_to(OUT)}")
    fixes = ", ".join(f"{a!r}->{b!r}:{n}" for (a, b), n in fixer.stats.most_common(12))
    print("glyph fixes:", fixes)
    untrusted = sorted({f"{d}:{f}" for (d, f), ok in fixer.trusted.items() if not ok})
    if untrusted:
        print("untrusted Sarabun fonts:", *untrusted, sep="\n  ")
    if not only:
        write_quality()
    return 0


def write_quality() -> None:
    """Score every converted file so a reader knows how far to trust it.

    ``unknown`` is the share of Thai words (two letters or more) missing from
    the PyThaiNLP dictionary. Names and loanwords keep it above zero even
    for clean text; broken decoding shows up as tens of percent.
    """
    from pythainlp.corpus import thai_words
    from pythainlp.tokenize import word_tokenize

    words = thai_words()
    rows = []
    for md in sorted(OUT.rglob("*.md")):
        if md.parent == OUT:
            continue
        text = md.read_text(encoding="utf-8")
        tokens = [w for run in re.findall(r"[ก-๎]+", text)
                  for w in word_tokenize(run, engine="newmm") if len(w) > 1]
        unknown = sum(w not in words for w in tokens) / max(len(tokens), 1)
        stray = len(re.findall("[�-]", text))
        odd = len(re.findall(
            r"(?<=[ก-๎])[^฀-๿\x00-\x7f -⁯ ✧•●○◆■□▪➢✓×÷±≤≥≠−…]"
            r"(?=[ก-๎])", text))
        pages = text.count("<!-- page ")
        rows.append(f"| {md.relative_to(OUT).as_posix()} | {pages} | {unknown:.1%} | {stray} | {odd} |")
    body = (
        "# Conversion quality\n\n"
        "Generated by `scripts/curriculum_sources_to_md.py`.\n\n"
        "- **Unknown words**: share of Thai words not in the PyThaiNLP dictionary.\n"
        "  Clean files sit near 0–2%; the rest is names and loanwords (เวเลนต์, ผศ.).\n"
        "- **Stray glyphs**: U+FFFD or private-use characters left in the text.\n"
        "- **Odd in Thai**: non-Thai characters sandwiched between Thai letters,\n"
        "  the usual sign of a glyph decoded to the wrong character.\n\n"
        "Formulas set as fractions, vectors or subscripts come out flattened in\n"
        "every file (`v = s t`). Check any formula against the PDF page.\n\n"
        "| File | Pages | Unknown words | Stray glyphs | Odd in Thai |\n"
        "|---|---:|---:|---:|---:|\n" + "\n".join(rows) + "\n"
    )
    (OUT / "QUALITY.md").write_text(body, encoding="utf-8", newline="\n")


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
